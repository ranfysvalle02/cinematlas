"""Loaders: turn files into records a collection can add.

A loader is any iterable of dicts. Subclass :class:`Loader` to get a suggested ``embed`` spec and
``moment`` field, so ``atlas.collection("decks", like=PDFPages)`` needs no further configuration.
"""

from __future__ import annotations

import json
from collections.abc import Iterator
from pathlib import Path
from typing import Any

from ..exceptions import DependencyError
from .parts import Image, Joint, Part, Text

IMAGE_SUFFIXES = {".jpg", ".jpeg", ".png", ".webp", ".gif", ".bmp", ".tif", ".tiff"}


class Loader:
    """Yields records. Class attributes suggest how to embed them."""

    #: Parts that make up the joint vector for these records.
    embed: Part | Joint | None = None
    #: Text field the reranker searches for the best sentence.
    moment: str | None = None
    #: Field that identifies a record, so re-adding replaces it.
    key: str | None = None
    #: Fields worth filtering on.
    filters: tuple[str, ...] = ()

    def __iter__(self) -> Iterator[dict[str, Any]]:
        raise NotImplementedError


class PDFPages(Loader):
    """One record per page: the rendered page image and its text. Slides, papers, manuals.

    Needs PyMuPDF: ``pip install 'cinematlas[pdf]'``.
    """

    embed = Text("text") + Image("image")
    moment = "text"
    key = "id"
    filters = ("source",)

    def __init__(self, path: str | Path, *, dpi: int = 110):
        self.path = Path(path)
        self.dpi = dpi

    def __iter__(self) -> Iterator[dict[str, Any]]:
        try:
            import pymupdf
        except ImportError as e:
            raise DependencyError("PDFPages needs PyMuPDF: pip install 'cinematlas[pdf]'") from e
        from PIL import Image as PILImage

        with pymupdf.open(self.path) as doc:
            for number, page in enumerate(doc, start=1):
                pix = page.get_pixmap(dpi=self.dpi)
                image = PILImage.frombytes("RGB", (pix.width, pix.height), pix.samples)
                yield {"id": f"{self.path.name}#{number}", "source": self.path.name, "page": number,
                       "text": page.get_text().strip(), "image": image}


class ImageFolder(Loader):
    """One record per image file, with an optional caption from a same-named ``.txt`` file."""

    embed = Image("path") + Text("caption")
    moment = "caption"
    key = "path"
    filters = ("folder",)

    def __init__(self, folder: str | Path, *, recursive: bool = True):
        self.folder = Path(folder)
        self.recursive = recursive

    def __iter__(self) -> Iterator[dict[str, Any]]:
        files = self.folder.rglob("*") if self.recursive else self.folder.glob("*")
        for path in sorted(files):
            if path.suffix.lower() not in IMAGE_SUFFIXES:
                continue
            caption_file = path.with_suffix(".txt")
            yield {"path": str(path), "name": path.stem, "folder": path.parent.name,
                   "caption": caption_file.read_text().strip() if caption_file.is_file() else None}


class JSONLines(Loader):
    """One record per line of a ``.jsonl`` file. Pair it with your own ``embed`` spec."""

    def __init__(self, path: str | Path):
        self.path = Path(path)

    def __iter__(self) -> Iterator[dict[str, Any]]:
        with self.path.open() as f:
            for line in f:
                if line.strip():
                    yield json.loads(line)


class Slides(Loader):
    """One record per slide: title, body text and **speaker notes** from a ``.pptx``.

    Pass the deck's PDF export (``pdf=``) to add each slide's rendered image, so the joint vector sees
    the slide itself too: charts, diagrams, screenshots. Needs python-pptx: ``pip install 'cinematlas[slides]'``.

        decks = atlas.collection("decks", like=Slides)
        decks.add(Slides("q3-review.pptx", pdf="q3-review.pdf"))
        decks.search("the slide where we showed Q3 churn")
    """

    embed = Text("title", label="Slide") + Text("body") + Text("notes", label="Speaker notes") + Image("image")
    moment = "notes_and_body"
    key = "id"
    filters = ("deck",)

    def __init__(self, path: str | Path, *, pdf: str | Path | None = None, dpi: int = 90):
        self.path = Path(path)
        self.pdf = Path(pdf) if pdf else None
        self.dpi = dpi

    def __iter__(self) -> Iterator[dict[str, Any]]:
        try:
            from pptx import Presentation
        except ImportError as e:
            raise DependencyError("Slides needs python-pptx: pip install 'cinematlas[slides]'") from e
        images = [page["image"] for page in PDFPages(self.pdf, dpi=self.dpi)] if self.pdf else []
        for number, slide in enumerate(Presentation(str(self.path)).slides, start=1):
            heading = slide.shapes.title
            heading_id = heading.shape_id if heading is not None else None  # proxies differ; ids don't
            texts = [(shape.shape_id == heading_id, shape.text_frame.text.strip()) for shape in slide.shapes
                     if shape.has_text_frame and shape.text_frame.text.strip()]
            title = next((t for is_title, t in texts if is_title), "")
            body = [t for is_title, t in texts if not is_title]
            notes = slide.notes_slide.notes_text_frame.text.strip() if slide.has_notes_slide else ""
            yield {"id": f"{self.path.name}#{number}", "deck": self.path.name, "slide": number,
                   "title": title, "body": "\n".join(body), "notes": notes,
                   "notes_and_body": "\n".join(filter(None, [notes, *body])),
                   "image": images[number - 1] if number <= len(images) else None}


class Screenshots(Loader):
    """One record per screenshot, with its on-screen text read by OCR, one line per visual line.

    The screenshot itself goes into the joint vector ("the screen with the red error banner"); the OCR
    text lets the reranker point at the matching line and lets you filter or display it. OCR needs
    ``pip install 'cinematlas[ocr]'``; without it, records carry the image alone.
    """

    embed = Image("path") + Text("text", label="On screen")
    moment = "text"
    key = "path"
    filters = ("folder",)

    def __init__(self, folder: str | Path, *, ocr: bool = True, recursive: bool = True):
        self.folder = Path(folder)
        self.ocr = ocr
        self.recursive = recursive

    def __iter__(self) -> Iterator[dict[str, Any]]:
        reader = _ocr_reader() if self.ocr else None
        try:
            for record in ImageFolder(self.folder, recursive=self.recursive):
                yield {"path": record["path"], "name": record["name"], "folder": record["folder"],
                       "text": _read_lines(reader, record["path"]) if reader else None}
        finally:
            # Release the ONNX session now, not at interpreter exit, where its teardown can race.
            del reader
            import gc

            gc.collect()


def _ocr_reader() -> Any:
    try:
        from rapidocr_onnxruntime import RapidOCR
    except ImportError:
        import logging

        logging.getLogger("cinematlas").warning(
            "Screenshots: OCR unavailable (pip install 'cinematlas[ocr]'); embedding images without their text.")
        return None
    return RapidOCR()


def _read_lines(reader: Any, path: str) -> str:
    """OCR fragments regrouped into visual lines, top to bottom, left to right."""
    import numpy as np
    from PIL import Image as PILImage

    result, _ = reader(np.array(PILImage.open(path).convert("RGB")))
    boxes = []
    for box, text, _score in result or []:
        ys, xs = [p[1] for p in box], [p[0] for p in box]
        boxes.append(((min(ys) + max(ys)) / 2, max(ys) - min(ys), min(xs), text))
    lines: list[list[tuple[float, float, float, str]]] = []
    for item in sorted(boxes):
        if lines and abs(item[0] - lines[-1][0][0]) < max(item[1], lines[-1][0][1]) / 2:
            lines[-1].append(item)
        else:
            lines.append([item])
    return "\n".join(" ".join(t for *_, t in sorted(line, key=lambda b: b[2])) for line in lines)
